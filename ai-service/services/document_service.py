"""Text extraction for PDF, Word, Excel, and CSV documents."""

import csv
import io
import logging
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from services.pdf_service import PageText, extract_pages

logger = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv"}


def extract_document(data: bytes, filename: str) -> tuple[int, list[PageText]]:
    """
    Extract text from a supported document.

    Returns:
        A tuple containing:
        - Number of pages, sheets, or logical sections.
        - Extracted text represented as PageText objects.
    """
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file format. Upload PDF, DOCX, XLSX, or CSV."
        )

    if extension == ".pdf":
        parsed = extract_pages(data)
        return parsed.page_count, parsed.pages

    if extension == ".docx":
        return _extract_docx(data)

    if extension == ".xlsx":
        return _extract_xlsx(data)

    if extension == ".csv":
        return _extract_csv(data)


def _extract_docx(data: bytes) -> tuple[int, list[PageText]]:
    """Extract paragraphs and table content from a Word document."""
    try:
        document = Document(io.BytesIO(data))

        lines = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table_number, table in enumerate(document.tables, start=1):
            lines.append(f"Table {table_number}:")

            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                lines.append(" | ".join(cells))

        text = "\n".join(lines).strip()

        if not text:
            raise ValueError("No extractable text found in the Word document.")

        return 1, [PageText(page=1, text=text)]

    except ValueError:
        raise
    except Exception as exc:
        logger.warning("Word extraction failed: %s", exc)
        raise ValueError("The Word document is corrupt or unreadable.") from exc


def _extract_xlsx(data: bytes) -> tuple[int, list[PageText]]:
    """Extract cell values from every worksheet in an Excel workbook."""
    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )

        pages = []

        try:
            for sheet_number, worksheet in enumerate(workbook.worksheets, start=1):
                lines = []

                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        str(value).strip() if value is not None else ""
                        for value in row
                    ]

                    if any(values):
                        lines.append(" | ".join(values))

                text = "\n".join(lines).strip()

                if text:
                    pages.append(
                        PageText(page=sheet_number, text=f"Sheet: {worksheet.title}\n{text}")
                    )
        finally:
            workbook.close()

        if not pages:
            raise ValueError("No extractable cell data found in the Excel workbook.")

        return len(workbook.sheetnames), pages

    except ValueError:
        raise
    except Exception as exc:
        logger.warning("Excel extraction failed: %s", exc)
        raise ValueError("The Excel workbook is corrupt or unreadable.") from exc


def _extract_csv(data: bytes) -> tuple[int, list[PageText]]:
    """Extract CSV rows and preserve their column structure."""
    try:
        text = data.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))

        lines = [
            " | ".join(cell.strip() for cell in row)
            for row in reader
            if any(cell.strip() for cell in row)
        ]

        extracted_text = "\n".join(lines).strip()

        if not extracted_text:
            raise ValueError("No extractable data found in the CSV file.")

        return 1, [PageText(page=1, text=extracted_text)]

    except UnicodeDecodeError as exc:
        raise ValueError(
            "The CSV encoding is unsupported. Please save it as UTF-8."
        ) from exc
    except csv.Error as exc:
        logger.warning("CSV extraction failed: %s", exc)
        raise ValueError("The CSV file is malformed or unreadable.") from exc